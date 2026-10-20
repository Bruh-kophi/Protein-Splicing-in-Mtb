"""fill_toc.py -- give the thesis a table of contents that is already filled in.

format_thesis.py inserts a Word TOC field, but Word leaves it empty until the
reader right-clicks > Update Field. This script writes the entries (every
Heading 2 and Heading 3 after the TOC) and their page numbers into that field,
so the contents show as soon as the file is opened. The field itself is kept,
so Word can still refresh it later (right-click > Update Field, or F9).

Page numbers are read from a PDF of the same document rendered by LibreOffice.
Word paginates very slightly differently, so if a heading falls near a page
end, Word's own update may differ by one page.

RUN  python3 fill_toc.py thesis.docx out.docx
     (needs LibreOffice `soffice` and poppler `pdftotext` on the PATH)
"""
import sys, os, re, subprocess, tempfile, copy
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]


def field_paragraph(d):
    for p in d.paragraphs:
        if any((e.text or "").strip().startswith("TOC") for e in p._p.iter(qn("w:instrText"))):
            return p
    sys.exit("no TOC field found")


def headings(d, toc_p):
    after, out_ = False, []
    for p in d.paragraphs:
        if p._p is toc_p._p:
            after = True; continue
        if after and p.style.name in ("Heading 2", "Heading 3") and p.text.strip():
            out_.append((int(p.style.name[-1]), p.text.strip()))
    return out_


def build(src_path, out_path, pages):
    d = docx.Document(src_path)
    toc_p = field_paragraph(d)
    entries = headings(d, toc_p)
    sec = d.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    # remove everything between "separate" and "end" in the field paragraph,
    # and any entry paragraphs written by a previous run of this script
    runs = list(toc_p._p.iter(qn("w:r")))
    state = 0; end_run = None
    for r in runs:
        types = [fc.get(qn("w:fldCharType")) for fc in r.findall(qn("w:fldChar"))]
        t = "end" if "end" in types else ("separate" if "separate" in types else None)
        if t == "separate":
            state = 1; continue
        if t == "end":
            end_run = r; break
        if state == 1:
            r.getparent().remove(r)
    end_run.getparent().remove(end_run)
    nxt = toc_p._p.getnext()
    while nxt is not None and nxt.find(qn("w:pPr")) is not None and \
            nxt.find(qn("w:pPr")).find(qn("w:pStyle")) is not None and \
            nxt.find(qn("w:pPr")).find(qn("w:pStyle")).get(qn("w:val")) == "TOCEntry":
        rm = nxt; nxt = nxt.getnext(); rm.getparent().remove(rm)
    # a paragraph style for entries (dot leader to a right tab at the margin)
    try:
        st = d.styles["TOCEntry"]
    except KeyError:
        st = d.styles.add_style("TOCEntry", 1)
        st.base_style = d.styles["Normal"]
    st.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.alignment = 0  # left
    anchor = toc_p._p
    first = True
    for (lvl, text), pg in zip(entries, pages + ["?"] * (len(entries) - len(pages))):
        if first:
            para = toc_p
            para.style = d.styles["TOCEntry"]; first = False
        else:
            el = OxmlElement("w:p"); anchor.addnext(el); anchor = el
            para = docx.text.paragraph.Paragraph(el, toc_p._parent)
            para.style = d.styles["TOCEntry"]
        para.paragraph_format.left_indent = Cm(0 if lvl == 2 else 0.75)
        # scientific names stay italic in the contents, as in the headings
        for k, piece in enumerate(re.split(r"(M\. tuberculosis|In silico)", text) + ["\t" + str(pg)]):
            if not piece:
                continue
            run = para.add_run(piece)
            run.font.name = "Times New Roman"; run.font.size = Pt(12)
            run.bold = (lvl == 2)
            run.italic = piece in ("M. tuberculosis", "In silico")
    # close the field after the last entry
    last = docx.text.paragraph.Paragraph(anchor, toc_p._parent)
    r = last.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
    d.save(out_path)
    return entries


def page_numbers(path, entries):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    page_lines = [[re.sub(r"\s+", " ", l).strip() for l in pg.splitlines()] for pg in txt.split("\f")]
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    # skip the TOC pages themselves: search from the page after the last TOC page
    start = next(i for i, pg in enumerate(page_lines) if "Table of Contents" in pg) + 1
    while start < len(page_lines) and any(l.endswith(tuple("0123456789")) and "Introduction" in l for l in page_lines[start]):
        start += 1
    pages, cur = [], start
    for _, text in entries:
        h = norm(text)
        found = None
        for i in range(cur, len(page_lines)):
            if any(l == h or l.startswith(h) for l in page_lines[i]):
                found = i; break
        if found is None:
            pages.append("?")
        else:
            pages.append(str(found + 1)); cur = found
    return pages


tmp_out = out + ".tmp.docx"
ents = build(src, tmp_out, [])
pages = page_numbers(tmp_out, ents)
for _ in range(2):  # repeat: filling in numbers can move pages
    build(src, tmp_out, pages)
    new = page_numbers(tmp_out, ents)
    if new == pages:
        break
    pages = new
build(src, out, pages)
os.remove(tmp_out)
for (lvl, t), p in zip(ents, pages):
    print(("  " if lvl == 3 else "") + t, "....", p)
print("saved", out)
